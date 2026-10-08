"""Build the Elevatus x Beyout AI certificate of attendance.

All text is live (Inter, embedded), so the PDF/AI opens in Illustrator
with editable type. Logos are the original vector paths.

  python3 build_certificate.py                      # the template (PDF + AI)
  python3 build_certificate.py --excel names.xlsx   # PDF + AI per name in column A, plus All_Certificates.pdf
"""
import argparse
import json
import os
import re
import shutil

from reportlab.lib.colors import HexColor
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")
FONTS = os.path.join(HERE, "fonts")

PAGE_W, PAGE_H = 792, 612  # US Letter landscape, same as the original artboard
CX = PAGE_W / 2

INK = HexColor("#11162B")
MUTED = HexColor("#586177")
NAVY = HexColor("#1E2B5A")      # Beyout AI navy, used in the seal
HAIRLINE = HexColor("#B9C0D0")
LOGO_SCALE = 0.85

# ---- Content (edit here, or in Illustrator) --------------------------------
CONTENT = {
    "title": "CERTIFICATE OF ATTENDANCE",
    "certifies": "THIS CERTIFIES THAT",
    "name": "Amr Khalid",
    "participated": "has successfully participated in the session",
    "session": "EMPLOYEE MAPPING & COMPETENCY FRAMEWORK POWERED BY AGENTIC AI",
    "description": [
        "Presented by Elevatus in partnership with Beyout AI, exploring how Agentic AI enables employee mapping,",
        "competency frameworks, workforce readiness, and data-driven talent decisions.",
    ],
    "awarded": "Awarded on 15 October 2026",
    "seal": ["EMPLOYEE MAPPING", "CERTIFICATE"],
    "footer": "ELEVATUS  ×  Beyout AI     |     elevatus.io",
}


def register_fonts():
    for w in ("Regular", "Medium", "SemiBold"):
        pdfmetrics.registerFont(TTFont(f"Inter-{w}", os.path.join(FONTS, f"Inter-{w}.ttf")))


def text(c, s, y, font, size, color=INK, tracking=0.0, x=CX):
    """One centred point-text line. `tracking` is in 1/1000 em, like Illustrator."""
    cs = size * tracking / 1000.0
    w = pdfmetrics.stringWidth(s, font, size) + cs * (len(s) - 1)
    t = c.beginText(x - w / 2, PAGE_H - y)
    t.setFont(font, size)
    t.setCharSpace(cs)
    t.setFillColor(color)
    t.textOut(s)
    c.drawText(t)


def fit_size(s, font, size, max_w):
    w = pdfmetrics.stringWidth(s, font, size)
    return size if w <= max_w else size * max_w / w


def draw_logos(c):
    """Original vector logo paths, lifted from the source artwork."""
    paths = json.load(open(os.path.join(ASSETS, "logos.json")))
    # Scale the lock-up about its centre (1.0 = original size).
    k, cx, cy = LOGO_SCALE, 396.0, 522.5
    ops = ["q", "%g 0 0 %g %g %g cm" % (k, k, cx * (1 - k), cy * (1 - k))]
    for p in paths:
        ops.append("q")
        ops.append("%g %g %g %g %g %g cm" % tuple(p["ctm"]))
        ops.append("%g %g %g rg" % tuple(p["fill"]))
        for op, v in p["segs"]:
            ops.append(" ".join("%g" % n for n in v) + (" " if v else "") + op)
        ops.append(p["op"])
        ops.append("Q")
    ops.append("Q")
    c._code.append("\n".join(ops))


def build(out_pdf, content):
    c = canvas.Canvas(out_pdf, pagesize=(PAGE_W, PAGE_H), initialFontName="Inter-Regular")
    c.setTitle("Certificate of Attendance")
    c.setAuthor("Elevatus")

    # Background field — same placement as the original artwork.
    c.drawImage(os.path.join(ASSETS, "background.jpg"),
                -101.894, -2.692, width=980.737, height=618.75)

    draw_logos(c)

    # Type stack. y = baseline, measured from the top of the page (pt).
    text(c, content["title"], 158, "Inter-Regular", 11, tracking=290)
    text(c, content["certifies"], 191, "Inter-Regular", 12, tracking=20)

    name_size = fit_size(content["name"], "Inter-SemiBold", 38, 620)
    text(c, content["name"], 247, "Inter-SemiBold", name_size, color=INK)
    c.setStrokeColor(HAIRLINE)
    c.setLineWidth(0.6)
    c.line(CX - 130, PAGE_H - 262, CX + 130, PAGE_H - 262)

    text(c, content["participated"], 284, "Inter-Regular", 10)
    text(c, content["session"], 318,
         "Inter-Medium", fit_size(content["session"], "Inter-Medium", 16, 660))
    for i, line in enumerate(content["description"]):
        text(c, line, 352 + i * 13, "Inter-Regular", 10, color=MUTED)
    text(c, content["awarded"], 410, "Inter-Regular", 12.5)

    # Seal: laurel artwork + live text in its centre.
    c.drawImage(ImageReader(os.path.join(ASSETS, "seal.png")),
                342.521, PAGE_H - 562.024, width=106.957, height=79.5, mask="auto")
    for i, line in enumerate(content["seal"]):
        text(c, line, 513 + i * 7.5, "Inter-Regular", 5.6, color=NAVY, tracking=20)

    text(c, content["footer"], 592, "Inter-Medium", 7.5, color=MUTED, tracking=120)

    c.showPage()
    c.save()


def safe(name):
    return re.sub(r"[^\w\- ]+", "", name).strip().replace(" ", "_") or "certificate"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--excel", help="xlsx with names in column A (row 1 = header)")
    ap.add_argument("--out", default=os.path.join(HERE, "output"))
    a = ap.parse_args()
    register_fonts()
    os.makedirs(a.out, exist_ok=True)

    if not a.excel:
        pdf = os.path.join(a.out, "Certificate_Template.pdf")
        build(pdf, CONTENT)
        shutil.copyfile(pdf, os.path.join(a.out, "Certificate_Template.ai"))
        print("wrote", pdf, "and .ai")
        return

    from openpyxl import load_workbook
    ws = load_workbook(a.excel, read_only=True).active
    from pypdf import PdfWriter
    pdf_dir, ai_dir = os.path.join(a.out, "PDF"), os.path.join(a.out, "AI")
    os.makedirs(pdf_dir, exist_ok=True)
    os.makedirs(ai_dir, exist_ok=True)
    merged, seen = PdfWriter(), {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None or not str(row[0]).strip():
            continue
        name = " ".join(str(row[0]).split())
        stem = f"Certificate_{safe(name)}"
        seen[stem] = seen.get(stem, 0) + 1
        if seen[stem] > 1:  # same name twice: keep both files
            stem += f"_{seen[stem]}"
        pdf = os.path.join(pdf_dir, stem + ".pdf")
        build(pdf, dict(CONTENT, name=name))
        shutil.copyfile(pdf, os.path.join(ai_dir, stem + ".ai"))
        merged.append(pdf)
    merged.write(os.path.join(a.out, "All_Certificates.pdf"))
    print(f"wrote {sum(seen.values())} certificates to {a.out}")


if __name__ == "__main__":
    main()
