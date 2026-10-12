# Elevatus Excel sheets (openpyxl): one function per building block, colours from brand/tokens.
# Inter is referenced by name: Excel shows it when the font is installed, otherwise it falls back to the system sans.
# Usage: see .claude/skills/elevatus-documents/SKILL.md
import os
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, 'assets')
LOGO, VIVID, ELECTRIC, APRICOT = '051274', '1B15CF', '2962FF', 'F0A15C'
INK, SLATE, LINE, PAPER, V100, WHITE = '11162B', '586177', 'E3E7F0', 'F4F7FF', 'E4EBFF', 'FFFFFF'

def font(size=10, bold=False, color=INK): return Font(name='Inter', size=size, bold=bold, color=color)
def fill(hex_): return PatternFill('solid', start_color=hex_, end_color=hex_)
_thin = Side(style='thin', color=LINE)
RULE = Border(bottom=_thin)

def workbook():
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = 'Elevatus'
    return wb

def sheet(wb, name, title, columns, rows, subtitle='', note=''):
    """One branded table sheet.
    columns: [(header, width, kind)] with kind in 'text' | 'int' | 'num' | 'pct' | 'date' | 'money'
    rows: list of lists, same order as columns. Returns the worksheet."""
    ws = wb.create_sheet(name[:31])
    ws.sheet_view.showGridLines = False
    n = len(columns)
    ws.column_dimensions['A'].width = 2.5  # left gutter
    for i, (_, width, _) in enumerate(columns):
        ws.column_dimensions[get_column_letter(i + 2)].width = width

    # masthead: wordmark, title, subtitle
    ws.row_dimensions[1].height = 12
    ws.row_dimensions[2].height = 26
    img = XLImage(os.path.join(ASSETS, 'wordmark-vivid.png'))
    img.width, img.height = 138, 138 * img.height / img.width
    ws.add_image(img, 'B2')
    ws.row_dimensions[3].height = 14
    ws.row_dimensions[4].height = 30
    ws['B4'] = title; ws['B4'].font = font(18, True, LOGO); ws['B4'].alignment = Alignment(vertical='center')
    r = 5
    if subtitle:
        ws.cell(r, 2, subtitle).font = font(10, color=SLATE); r += 1
    r += 1

    # table header
    head = r
    ws.row_dimensions[head].height = 24
    for i, (label, _, kind) in enumerate(columns):
        c = ws.cell(head, i + 2, label)
        c.font = font(10, True, WHITE); c.fill = fill(LOGO)
        c.alignment = Alignment(horizontal='left' if kind == 'text' else 'right', vertical='center', indent=1)
    # body
    FORMATS = {'int': '#,##0', 'num': '#,##0.00', 'pct': '0%', 'date': 'd mmm yyyy', 'money': '#,##0.00'}
    for k, row in enumerate(rows):
        rr = head + 1 + k
        ws.row_dimensions[rr].height = 20
        for i, (_, _, kind) in enumerate(columns):
            c = ws.cell(rr, i + 2, row[i] if i < len(row) else None)
            c.font = font(10); c.border = RULE
            c.alignment = Alignment(horizontal='left' if kind == 'text' else 'right', vertical='center', indent=1, wrap_text=kind == 'text')
            if kind in FORMATS: c.number_format = FORMATS[kind]
            if k % 2: c.fill = fill(PAPER)
    last = head + len(rows)
    if rows: ws.auto_filter.ref = f'B{head}:{get_column_letter(n + 1)}{last}'
    ws.freeze_panes = ws.cell(head + 1, 2)
    if note:
        ws.cell(last + 2, 2, note).font = font(9, color=SLATE)

    # print: A4 landscape, fit to width, header row repeats, branded footer
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f'{head}:{head}'
    ws.oddFooter.left.text = 'Elevatus'; ws.oddFooter.right.text = 'Page &P of &N'
    ws.sheet_properties.tabColor = VIVID
    return ws

def cover(wb, title, description='', category='', label=''):
    """Cover sheet in the same voice as the PDF cover: Vivid Blue field, white wordmark and type."""
    ws = wb.create_sheet('Cover', 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = LOGO
    for col in range(1, 13): ws.column_dimensions[get_column_letter(col)].width = 12
    for row in range(1, 34):
        for col in range(1, 13): ws.cell(row, col).fill = fill(VIVID)
    img = XLImage(os.path.join(ASSETS, 'wordmark-white.png'))
    img.width, img.height = 200, 200 * img.height / img.width
    ws.add_image(img, 'B3')
    ws.row_dimensions[8].height = 48
    ws['B8'] = title; ws['B8'].font = Font(name='Inter', size=34, color=WHITE); ws['B8'].alignment = Alignment(vertical='center')
    ws['B10'] = description; ws['B10'].font = font(12, color=WHITE)
    ws['B31'] = category.upper().replace('\n', ' '); ws['B31'].font = font(9, True, WHITE)
    ws['K31'] = label; ws['K31'].font = font(11, color=WHITE); ws['K31'].alignment = Alignment(horizontal='right')
    return ws
